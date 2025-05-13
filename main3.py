from agents import Agent, Runner, AsyncOpenAI, set_default_openai_api
from agents import set_default_openai_client, set_tracing_disabled
from agents.tool import function_tool
import os
from dotenv import load_dotenv

load_dotenv()

if not os.getenv("GEMINI_API_KEY"):
    raise ValueError("GEMINI_API_KEY is not set. Please ensure it is defined in your .env file.")

gemini_api_key = os.getenv('GEMINI_API_KEY')
set_tracing_disabled(False)

external_client = AsyncOpenAI(
    api_key=gemini_api_key,
    base_url="https://generativelanguage.googleapis.com/v1beta/openai/",
)
set_default_openai_client(external_client)
set_default_openai_api("chat_completions")
set_tracing_disabled(True)

# ----------------------------
# File Reader Tool Definition
# ----------------------------
import mimetypes
from pdfminer.high_level import extract_text
import openpyxl
import docx

@function_tool
def file_reader(file_path: str) -> str:
    """
    Reads content from a given file path and returns its textual content.
    Supports PDF, Excel, Word, and plain text files.
    """
    try:
        mime_type, _ = mimetypes.guess_type(file_path)

        if file_path.endswith(".pdf"):
            try:
                text = extract_text(file_path)
                if text.strip():
                    return text
                else:
                    return "The PDF was read, but no text was found. It may be scanned or image-based."
            except Exception as e:
                return f"An error occurred while reading the PDF: {str(e)}"

        elif file_path.endswith(".xlsx") or "spreadsheet" in str(mime_type):
            wb = openpyxl.load_workbook(file_path)
            content = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    content.append("\t".join([str(cell) if cell is not None else "" for cell in row]))
            return "\n".join(content) or "No readable text found in Excel."

        elif file_path.endswith(".docx") or "word" in str(mime_type):
            doc = docx.Document(file_path)
            return "\n".join([para.text for para in doc.paragraphs]) or "No readable text found in Word document."

        elif file_path.endswith(".txt") or "text" in str(mime_type):
            with open(file_path, "r", encoding="utf-8") as f:
                return f.read()

        else:
            return f"Unsupported file format: {file_path}"

    except Exception as e:
        return f"An error occurred while reading the file: {str(e)}"

# ----------------------------
# Agent Definition
# ----------------------------
agent = Agent(
    name="ResearcherBot",
    instructions=(
        "You are a research assistant. "
        "Use 'file_reader' to extract information from uploaded files. "
        "Do not summarize or modify tool outputs. Show maximum content."
    ),
    tools=[file_reader],
    model="gemini-2.0-flash",
)

# ----------------------------
# Usage Example
# ----------------------------
import os

# ✅ Use a TXT file (works fine) OR try a PDF
file_path = os.path.abspath("files\check.pdf")  # Try with file.txt too

query = f"The file is located here: {file_path}. Can you read and show me its content?"

result = Runner.run_sync(agent, query)

print("\n✅ Agent Output:\n")
print(result.final_output)
