#  Implement with Chainlit
#  This code is a Chainlit application that allows users to upload files and interact with an AI agent.
#  Agentic RAG is a research assistant that can read and analyze files.


import chainlit as cl
from agents import Agent, Runner, AsyncOpenAI, set_default_openai_api
from agents import set_default_openai_client, set_tracing_disabled
from agents.tool import function_tool
import os
from dotenv import load_dotenv
import mimetypes
from pdfminer.high_level import extract_text
import openpyxl
import docx

# Load environment variables
load_dotenv()

# Set up Gemini API
if not os.getenv("GEMINI_API_KEY"):
    raise ValueError("GEMINI_API_KEY is not set. Please ensure it is defined in your .env file.")
gemini_api_key = os.getenv('GEMINI_API_KEY')
set_tracing_disabled(False)
external_client = AsyncOpenAI(
    api_key=gemini_api_key,
    base_url="https://generativelanguage.googleapis.com/v1beta/openai/"
)
set_default_openai_client(external_client)
set_default_openai_api("chat_completions")
set_tracing_disabled(True)

# Define the file_reader tool
@function_tool
def file_reader(file_path: str) -> str:
    """
    Reads content from a given file path and returns its textual content.
    Supports PDF, Excel, Word, and plain text files.
    """
    try:
        mime_type, _ = mimetypes.guess_type(file_path)

        if file_path.endswith(".pdf"):
            text = extract_text(file_path)
            return text or "No readable text found in PDF."

        elif file_path.endswith(".xlsx") or "spreadsheet" in str(mime_type):
            wb = openpyxl.load_workbook(file_path)
            content = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    content.append("\t".join([str(cell) if cell is not None else "" for cell in row]))
            return "\n".join(content) or "No readable text found in Excel."

        elif file_path.endswith(".docx") or "word" in str(mime_type):
            doc = docx.Document(file_path)
            return "\n".join([para.text for para in doc.paragraphs]) or "No readable text found in Word document."

        elif file_path.endswith(".txt") or "text" in str(mime_type):
            with open(file_path, "r", encoding="utf-8") as f:
                return f.read()

        else:
            return f"Unsupported file format: {file_path}"

    except Exception as e:
        return f"An error occurred while reading the file: {str(e)}"

# Set up Agent
agent = Agent(
    name="ResearcherBot",
    instructions=(
        "You are a research assistant. "
        "Use 'file_reader' to extract content from files when the user mentions reading or analyzing a file. "
        "Always try to show the full content."
    ),
    tools=[file_reader],
    model="gemini-2.0-flash",
)

# Chainlit handlers
@cl.on_chat_start
async def start():
    # Display "Top AGENTIC RAG" as a heading at the top of the UI
    await cl.Message(content="# AGENTIC RAG").send()
    
    # Prompt user to upload a file
    files = await cl.AskFileMessage(
        content="Please upload a file to begin!",
        accept=["*/*"],  # Accept all file types
        max_size_mb=20,
        timeout=180,
    ).send()
    if files:
        file_path = files[0].path
        cl.user_session.set("file_path", file_path)
        await cl.Message(content=f"File uploaded: {files[0].name}").send()
    else:
        await cl.Message(content="No file uploaded. Please start over.").send()

@cl.on_message
async def main(message: cl.Message):
    file_path = cl.user_session.get("file_path")
    if not file_path:
        await cl.Message(content="Please upload a file first.").send()
        return

    user_query = message.content
    full_query = f"The file is located here: {file_path}. {user_query}"
    result = await cl.make_async(Runner.run_sync)(agent, full_query)
    await cl.Message(content=result.final_output).send()