import os
import mimetypes
from dotenv import load_dotenv
import chainlit as cl
from pdfminer.high_level import extract_text
import openpyxl
import docx

from agents import Agent, AsyncOpenAI, set_default_openai_client, set_default_openai_api, set_tracing_disabled
from agents.tool import function_tool

# ✅ Load environment variables
load_dotenv()
gemini_api_key = os.getenv("GEMINI_API_KEY")

if not gemini_api_key:
    raise ValueError("GEMINI_API_KEY is not set in .env")

# ✅ Set up Gemini OpenAI-style client
external_client = AsyncOpenAI(
    api_key=gemini_api_key,
    base_url="https://generativelanguage.googleapis.com/v1beta/openai/",
)
set_default_openai_client(external_client)
set_default_openai_api("chat_completions")
set_tracing_disabled(True)

# ✅ File Reader Tool
@function_tool
def file_reader(file_path: str) -> str:
    """
    Reads content from PDF, Word, Excel, or text file.
    """
    try:
        mime_type, _ = mimetypes.guess_type(file_path)

        if file_path.endswith(".pdf"):
            return extract_text(file_path) or "No readable text in PDF."

        elif file_path.endswith(".xlsx") or "spreadsheet" in str(mime_type):
            wb = openpyxl.load_workbook(file_path)
            content = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    content.append("\t".join([str(cell) if cell else "" for cell in row]))
            return "\n".join(content) or "No readable text in Excel."

        elif file_path.endswith(".docx") or "word" in str(mime_type):
            doc = docx.Document(file_path)
            return "\n".join([para.text for para in doc.paragraphs]) or "No readable text in Word document."

        elif file_path.endswith(".txt") or "text" in str(mime_type):
            with open(file_path, "r", encoding="utf-8") as f:
                return f.read()

        else:
            return f"Unsupported file format: {file_path}"

    except Exception as e:
        return f"Error reading file: {str(e)}"

# ✅ Create Agent (still useful for future messages)
agent = Agent(
    name="ResearcherBot",
    instructions=(
        "You are a research assistant. "
        "Use 'file_reader' to extract content from files when the user mentions reading or analyzing a file."
    ),
    tools=[file_reader],
    model="gemini-2.0-flash",
)

# ✅ Start Chat
@cl.on_chat_start
async def start_chat():
    cl.user_session.set("agent", agent)
    await cl.Message(content="📄 Hi! Upload a file or ask a question.").send()

# ✅ Handle messages and files
@cl.on_message
async def handle_message(message: cl.Message):
    agent = cl.user_session.get("agent")

    # Handle file uploads
    if message.elements:
        for element in message.elements:
            if isinstance(element, cl.File):
                file_path = f"uploaded_files/{element.name}"
                os.makedirs("uploaded_files", exist_ok=True)
                await element.save(file_path)

                # ✅ Read file content directly and respond
                file_text = file_reader(file_path)
                await cl.Message(content=file_text).send()
        return

    # Handle regular chat messages (not file uploads)
    response = await agent.run(message.content)
    await cl.Message(content=response).send()
